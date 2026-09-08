"""
Script de verificación manual programática para la exportación de resultados a Word (.docx).
Genera un documento Word con el texto no confidencial del ejemplo y valida
exhaustivamente mediante python-docx su estructura, estilos, párrafos, viñetas y negritas.
"""

import io
import sys
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))

import docx
from app.schemas.export import WordExportMetadata, WordExportRequest
from app.services.word_export_service import word_export_service

SAMPLE_TITLE = "Resumen Ejecutivo de Servicios"
SAMPLE_CONTENT = """# Resumen ejecutivo

El contrato fue firmado el 3 de marzo de 2026.

## Obligaciones principales

- Empresa Alfa debe entregar un informe mensual.
- El informe debe entregarse antes del día 5.

## Penalizaciones

Existe una penalización de **500 euros** por retrasos superiores a diez días."""

SAMPLE_WARNINGS = [
    "No se ha identificado cláusula de resolución anticipada en el documento de prueba."
]


def run_manual_verification():
    print("=== VERIFICACIÓN PROGRAMÁTICA: EXPORTACIÓN WORD (.DOCX) ===")

    request = WordExportRequest(
        title=SAMPLE_TITLE,
        content=SAMPLE_CONTENT,
        warnings=SAMPLE_WARNINGS,
        metadata=WordExportMetadata(
            prompt_name="Resumen ejecutivo",
            model="gemini-3.8-flash",
        ),
    )

    # 1. Generar DOCX en memoria
    print("[1/5] Generando documento Word en memoria (io.BytesIO)...")
    buffer = word_export_service.generate_docx(request)
    doc_bytes = buffer.getvalue()
    print(f"      Tamaño generado en RAM: {len(doc_bytes)} bytes")
    assert len(doc_bytes) > 0, "El búfer de bytes está vacío"

    # 2. Abrir con python-docx para validar estructura
    print("[2/5] Inspeccionando estructura interna con python-docx...")
    doc = docx.Document(io.BytesIO(doc_bytes))
    print(f"      Párrafos totales: {len(doc.paragraphs)}")
    assert len(doc.paragraphs) >= 8, f"Se esperaban al menos 8 párrafos, obtenidos {len(doc.paragraphs)}"

    # 3. Validar título y subtítulo de metadatos
    print("[3/5] Validando Título, Metadatos y Headings...")
    first_p = doc.paragraphs[0]
    print(f"      Título (style={first_p.style.name}): '{first_p.text}'")
    assert first_p.text == SAMPLE_TITLE
    assert first_p.style.name == "Title"

    meta_p = doc.paragraphs[1]
    print(f"      Subtítulo metadatos: '{meta_p.text}'")
    assert "Tipo de análisis: Resumen ejecutivo" in meta_p.text
    assert "gemini-3.8-flash" not in meta_p.text  # Confirmar que el modelo no se expone al cliente

    headings = [(p.style.name, p.text) for p in doc.paragraphs if "Heading" in p.style.name]
    print(f"      Headings detectados: {headings}")
    assert any(h[0] == "Heading 1" and h[1] == "Resumen ejecutivo" for h in headings)
    assert any(h[0] == "Heading 2" and h[1] == "Obligaciones principales" for h in headings)
    assert any(h[0] == "Heading 2" and h[1] == "Penalizaciones" for h in headings)
    assert any(h[0] == "Heading 2" and h[1] == "Advertencias" for h in headings)

    # 4. Validar viñetas, negrita y texto
    print("[4/5] Validando listas de viñetas, párrafos y formato inline (bold)...")
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    print(f"      Viñetas encontradas ({len(bullets)}):")
    for b in bullets:
        print(f"        * {b}")
    assert any("Empresa Alfa debe entregar un informe mensual." in b for b in bullets)
    assert any("El informe debe entregarse antes del día 5." in b for b in bullets)
    assert any("No se ha identificado cláusula de resolución anticipada" in b for b in bullets)

    bold_runs = [run.text for p in doc.paragraphs for run in p.runs if run.bold]
    print(f"      Segmentos en negrita encontrados: {bold_runs}")
    assert "500 euros" in bold_runs, "No se encontró el formato inline de negrita para '500 euros'"

    # 5. Validar pie de página
    print("[5/5] Validando pie de página de sección...")
    footer_text = doc.sections[0].footer.paragraphs[0].text
    print(f"      Pie de página: '{footer_text}'")
    assert "Generado mediante HITCHINGS" in footer_text

    # 6. Validar nombre de archivo
    filename = word_export_service.sanitize_filename(request.title)
    print(f"      Nombre de archivo sanitizado: '{filename}'")
    assert filename == "resumen-ejecutivo-de-servicios.docx"

    print("\n===========================================================")
    print(">>> VERIFICACIÓN MANUAL PROGRAMÁTICA: PASS <<<")
    print("===========================================================")


if __name__ == "__main__":
    run_manual_verification()
