import io
import logging
from unittest.mock import patch
import docx
import pytest
from fastapi.testclient import TestClient

from app.core.config import settings
from app.services.word_export_service import word_export_service


SAMPLE_EXPORT_PAYLOAD = {
    "title": "Análisis Jurídico del Contrato",
    "content": (
        "# Resumen ejecutivo\n\n"
        "El contrato fue suscrito con fecha 3 de marzo de 2026 entre las partes intervinientes.\n\n"
        "## Obligaciones principales\n\n"
        "- Empresa Alfa debe entregar un informe mensual antes del día 5.\n"
        "* Empresa Beta debe validar la conformidad en un plazo de diez días.\n\n"
        "## Fases del procedimiento\n\n"
        "1. Notificación formal mediante burofax.\n"
        "2. Interposición de demanda ejecutiva.\n\n"
        "## Penalizaciones\n\n"
        "Existe una penalización de **500 euros** por retrasos de *diez días* hábiles.\n\n"
        "| Concepto | Importe |\n"
        "| --- | --- |\n"
        "| Cuota fija | 1.000 € |\n"
        "| Penalización | 500 € |\n"
    ),
    "warnings": [
        "No se ha identificado la cláusula de sumisión a arbitraje.",
        "Falta firma de la parte compradora en el anexo II.",
    ],
    "metadata": {
        "prompt_name": "Análisis jurídico",
        "model": "gemini-3.8-flash",
    },
}


# 1. Exportación válida (HTTP 200)
def test_export_word_valid(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    assert len(response.content) > 0


# 2. Content-Type correcto
def test_export_word_content_type(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    expected_content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert response.headers["content-type"] == expected_content_type


# 3. Content-Disposition con attachment y .docx
def test_export_word_content_disposition(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    cd = response.headers.get("content-disposition", "")
    assert "attachment" in cd
    assert 'filename="analisis-juridico-del-contrato.docx"' in cd


# 4. Fichero resultante abre correctamente con python-docx
def test_export_word_opens_validly_with_python_docx(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))
    assert doc is not None
    assert len(doc.paragraphs) > 5


# 5. Título principal aparece en el Word
def test_export_word_title_rendered(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Análisis Jurídico del Contrato" in all_text

    # Comprobar que el primer párrafo tiene estilo Title
    first_p = doc.paragraphs[0]
    assert first_p.text == "Análisis Jurídico del Contrato"
    assert first_p.style.name == "Title"


# 6. Párrafos normales aparecen en el Word
def test_export_word_paragraphs_rendered(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "El contrato fue suscrito con fecha 3 de marzo de 2026" in all_text


# 7. Headings Markdown (#, ##, ###) convertidos a Heading 1, 2, 3
def test_export_word_headings_converted(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    heading_styles = [p.style.name for p in doc.paragraphs if "Heading" in p.style.name]
    assert "Heading 1" in heading_styles
    assert "Heading 2" in heading_styles

    h1_texts = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert "Resumen ejecutivo" in h1_texts

    h2_texts = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert "Obligaciones principales" in h2_texts


# 8. Listas de viñetas (- y *) convertidas a List Bullet
def test_export_word_bullets_converted(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    bullet_paragraphs = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert len(bullet_paragraphs) >= 2
    bullet_texts = [p.text for p in bullet_paragraphs]
    assert any("Empresa Alfa debe entregar" in t for t in bullet_texts)
    assert any("Empresa Beta debe validar" in t for t in bullet_texts)


# 9. Listas numeradas convertidas a List Number
def test_export_word_numbered_list_converted(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    numbered_paragraphs = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert len(numbered_paragraphs) >= 2
    numbered_texts = [p.text for p in numbered_paragraphs]
    assert any("Notificación formal" in t for t in numbered_texts)
    assert any("Interposición de demanda" in t for t in numbered_texts)


# 10. Negrita básica aplicada a runs
def test_export_word_bold_run_applied(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    bold_runs = [run for p in doc.paragraphs for run in p.runs if run.bold]
    bold_texts = [r.text for r in bold_runs]
    assert any("500 euros" in t for t in bold_texts)


# 11. Cursiva básica aplicada a runs
def test_export_word_italic_run_applied(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    italic_runs = [run for p in doc.paragraphs for run in p.runs if run.italic]
    italic_texts = [r.text for r in italic_runs]
    assert any("diez días" in t for t in italic_texts)


# 12. Sección de advertencias aparece cuando warnings != []
def test_export_word_warnings_rendered_when_present(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Advertencias" in all_text
    assert "No se ha identificado la cláusula de sumisión a arbitraje" in all_text
    assert "Falta firma de la parte compradora" in all_text


# 13. Warnings vacíos no generan sección de advertencias
def test_export_word_no_warnings_section_when_empty(client: TestClient):
    payload = {
        "title": "Análisis Jurídico Limpio",
        "content": "## Contenido\nTodo correcto sin incidencias detectadas.",
        "warnings": [],
    }
    response = client.post("/api/v1/export/word", json=payload)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Advertencias" not in all_text


# 14. Metadata discreta en el documento
def test_export_word_metadata_rendered(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Tipo de análisis: Análisis jurídico" in all_text
    # El modelo de IA se preserva en metadata de API pero no ensucia el cuerpo
    assert "gemini-3.8-flash" not in all_text


def test_export_word_renders_prompt_name_but_never_model(client: TestClient):
    payload = {
        "title": "Dictamen de Prueba",
        "content": "Contenido del dictamen legal.",
        "warnings": [],
        "metadata": {
            "prompt_name": "Puntos clave",
            "model": "gemini-3.8-flash",
        },
    }
    response = client.post("/api/v1/export/word", json=payload)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))
    all_text = " ".join(p.text for p in doc.paragraphs)

    # El nombre de la plantilla sí debe aparecer visiblemente
    assert "Tipo de análisis: Puntos clave" in all_text
    # El modelo interno bajo ningún concepto debe aparecer en el Word
    assert "gemini-3.8-flash" not in all_text
    assert "gemini" not in all_text.lower()



# 15. Sanitización de nombre de archivo seguro
def test_export_word_filename_sanitization():
    # Títulos con acentos, caracteres extraños y espacios
    assert word_export_service.sanitize_filename("Sentencia Estimatoria Núm. 45/2026") == "sentencia-estimatoria-num-45-2026.docx"
    assert word_export_service.sanitize_filename("  Análisis & Dictamen !!!  ") == "analisis-dictamen.docx"
    assert word_export_service.sanitize_filename("$$$###%%%") == "hitchings-analisis.docx"
    assert word_export_service.sanitize_filename("") == "hitchings-analisis.docx"


# 16. Título vacío devuelve HTTP 400
def test_export_word_empty_title_returns_400(client: TestClient):
    for empty_title in ["", "   ", "\n\t"]:
        payload = {
            "title": empty_title,
            "content": "Contenido válido para exportar.",
        }
        response = client.post("/api/v1/export/word", json=payload)
        # Puede ser 400 o 422 según validación Pydantic
        assert response.status_code in (400, 422)


# 17. Contenido vacío devuelve HTTP 400
def test_export_word_empty_content_returns_400(client: TestClient):
    for empty_content in ["", "   ", "\n\t\n"]:
        payload = {
            "title": "Título Válido",
            "content": empty_content,
        }
        response = client.post("/api/v1/export/word", json=payload)
        assert response.status_code in (400, 422)


# 18. Exceso de caracteres devuelve HTTP 413
def test_export_word_exceeds_max_characters_returns_413(client: TestClient):
    with patch.object(settings, "MAX_EXPORT_CHARACTERS", 100):
        payload = {
            "title": "Título Válido",
            "content": "X" * 101,
        }
        response = client.post("/api/v1/export/word", json=payload)
        assert response.status_code == 413
        assert "excede el límite máximo permitido" in response.json()["detail"]


# 19. Confidencialidad en logs (cero texto de contenido o título confidencial)
def test_export_word_confidentiality_logs(client: TestClient, caplog):
    secret_title = "CONFIDENTIAL_INSPECTION_AUDIT_REPORT_889900"
    secret_text = "SECRET_FINANCIAL_SETTLEMENT_AMOUNT_123456789"
    payload = {
        "title": secret_title,
        "content": f"El informe secreto dice: {secret_text}",
        "warnings": ["Advertencia confidencial interna 998877"],
    }
    with caplog.at_level(logging.INFO):
        response = client.post("/api/v1/export/word", json=payload)
        assert response.status_code == 200

    logs = caplog.text
    assert secret_title not in logs
    assert secret_text not in logs
    assert "Exportación a Word completada" in logs


# 20. El Word no se escribe permanentemente en disco
def test_export_word_memory_only(client: TestClient, tmp_path):
    # Verificar que tras ejecutar una exportación no se crean archivos residuales en el directorio actual
    import os

    files_before = set(os.listdir("."))
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    files_after = set(os.listdir("."))

    # No deben haber aparecido archivos .docx en el directorio de trabajo
    new_files = files_after - files_before
    assert not any(f.endswith(".docx") for f in new_files)


# 21. Tablas Markdown convertidas a tablas Word
def test_export_word_tables_converted(client: TestClient):
    response = client.post("/api/v1/export/word", json=SAMPLE_EXPORT_PAYLOAD)
    assert response.status_code == 200
    doc = docx.Document(io.BytesIO(response.content))

    assert len(doc.tables) >= 1
    table = doc.tables[0]
    # Comprobar cabecera
    assert table.cell(0, 0).text.strip() == "Concepto"
    assert table.cell(0, 1).text.strip() == "Importe"
    # Comprobar valores
    assert table.cell(1, 0).text.strip() == "Cuota fija"
    assert table.cell(2, 0).text.strip() == "Penalización"
