import io
import re
import unicodedata
from typing import List

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.core.config import settings
from app.core.logging import logger
from app.schemas.export import WordExportRequest


class WordExportError(Exception):
    """Excepción base para errores durante la exportación a Word."""
    pass


class WordExportEmptyContentError(WordExportError):
    """Excepción cuando el título o contenido a exportar está vacío."""
    pass


class WordExportSizeExceededError(WordExportError):
    """Excepción cuando el contenido excede el límite máximo permitido."""
    def __init__(self, char_count: int, max_chars: int):
        super().__init__(
            f"El contenido a exportar ({char_count:,} caracteres) excede el límite máximo permitido ({max_chars:,} caracteres)."
        )


class WordExportService:
    """Servicio para convertir resultados de análisis estructurados a documentos Microsoft Word (.docx)."""

    FONT_NAME = "Arial"
    DEFAULT_FILENAME = "hitchings-analisis.docx"

    def sanitize_filename(self, title: str) -> str:
        """
        Genera un nombre de archivo seguro para descarga a partir del título:
        - Normaliza tildes y caracteres especiales a ASCII
        - Convierte a minúsculas
        - Reemplaza espacios y signos no alfanuméricos por guiones
        - Evita secuencias de guiones repetidos
        - Limita la longitud a 60 caracteres
        - Asegura extensión .docx
        - Si el resultado es vacío, recurre a 'hitchings-analisis.docx'
        """
        if not title or not title.strip():
            return self.DEFAULT_FILENAME

        # 1. Normalizar acentos a ASCII
        normalized = (
            unicodedata.normalize("NFKD", title.strip())
            .encode("ascii", "ignore")
            .decode("ascii")
        )

        # 2. Minúsculas y reemplazar caracteres inválidos por guiones
        cleaned = re.sub(r"[^a-zA-Z0-9]+", "-", normalized.lower())

        # 3. Eliminar guiones iniciales, finales y duplicados
        cleaned = re.sub(r"-+", "-", cleaned).strip("-")

        # 4. Limitar longitud máxima
        if len(cleaned) > 60:
            cleaned = cleaned[:60].rstrip("-")

        if not cleaned:
            return self.DEFAULT_FILENAME

        return f"{cleaned}.docx"

    def _strip_html(self, text: str) -> str:
        """Elimina de forma segura cualquier etiqueta HTML embebida."""
        return re.sub(r"<[^>]+>", "", text)

    def _add_formatted_runs(self, paragraph, text: str) -> None:
        """
        Interpreta formatos inline sencillos de Markdown (**negrita**, *cursiva*)
        y los añade como runs formateados al párrafo suministrado.
        """
        clean_text = self._strip_html(text)
        if not clean_text:
            return

        # Patrón para identificar segmentos **negrita**, *cursiva* o texto normal
        # Evita confundir ** con * mediante orden de precedencia
        pattern = re.compile(r"(\*\*.*?\*\*|\*[^*]+?\*|[^*]+)")
        tokens = pattern.findall(clean_text)

        for token in tokens:
            if token.startswith("**") and token.endswith("**") and len(token) >= 4:
                run = paragraph.add_run(token[2:-2])
                run.bold = True
            elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            else:
                paragraph.add_run(token)

    def _parse_table_block(self, doc: docx.Document, table_lines: List[str]) -> None:
        """Convierte un bloque acumulado de líneas de tabla Markdown a una tabla nativa de Word."""
        data_rows: List[List[str]] = []

        for line in table_lines:
            raw_line = line.strip()
            # Ignorar fila divisoria Markdown (| --- | --- |)
            if re.match(r"^\|(\s*:?-+:?\s*\|)+$", raw_line):
                continue
            # Extraer celdas
            cells = [c.strip() for c in raw_line.strip("|").split("|")]
            data_rows.append(cells)

        if not data_rows:
            return

        # Normalizar número de columnas según la fila con más celdas
        col_count = max(len(row) for row in data_rows)
        if col_count == 0:
            return

        table = doc.add_table(rows=len(data_rows), cols=col_count)
        table.style = "Table Grid"

        for row_idx, row_data in enumerate(data_rows):
            for col_idx in range(col_count):
                cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
                cell = table.cell(row_idx, col_idx)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                self._add_formatted_runs(p, cell_text)
                if row_idx == 0:
                    for run in p.runs:
                        run.bold = True

        # Añadir un espacio después de la tabla
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(6)

    def _render_markdown_content(self, doc: docx.Document, content: str) -> None:
        """
        Interpreta el contenido Markdown generado y añade los elementos correspondientes
        con estilos nativos de Microsoft Word.
        """
        lines = content.splitlines()
        i = 0
        n = len(lines)

        while i < n:
            raw_line = lines[i]
            stripped = raw_line.strip()

            # 1. Líneas en blanco
            if not stripped:
                i += 1
                continue

            # 2. Separador horizontal (---, ***, ___)
            if re.match(r"^(---|[*]{3,}|_{3,})$", stripped):
                i += 1
                continue

            # 3. Tablas Markdown (acumular líneas consecutivas que comiencen por '|')
            if stripped.startswith("|") and stripped.endswith("|"):
                table_lines = []
                while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                self._parse_table_block(doc, table_lines)
                continue

            # 4. Encabezados (#, ##, ###, ####)
            heading_match = re.match(r"^(#{1,4})\s+(.+)$", stripped)
            if heading_match:
                level_str, heading_text = heading_match.groups()
                level = len(level_str)
                # Word admite headings de nivel 1 a 9
                h = doc.add_heading(level=level)
                self._add_formatted_runs(h, heading_text)
                i += 1
                continue

            # 5. Listas de viñetas (- elemento, * elemento, + elemento)
            bullet_match = re.match(r"^[-*+]\s+(.+)$", stripped)
            if bullet_match:
                item_text = bullet_match.group(1)
                p = doc.add_paragraph(style="List Bullet")
                self._add_formatted_runs(p, item_text)
                i += 1
                continue

            # 6. Listas numeradas (1. elemento, 2. elemento)
            number_match = re.match(r"^\d+\.\s+(.+)$", stripped)
            if number_match:
                item_text = number_match.group(1)
                p = doc.add_paragraph(style="List Number")
                self._add_formatted_runs(p, item_text)
                i += 1
                continue

            # 7. Párrafo estándar
            p = doc.add_paragraph(style="Normal")
            self._add_formatted_runs(p, stripped)
            i += 1

    def generate_docx(self, request: WordExportRequest) -> io.BytesIO:
        """
        Genera el documento Word en memoria aplicando las especificaciones de HITCHINGS:
        - Título principal
        - Subtítulo discreto de metadatos (si existe)
        - Contenido principal interpretado desde Markdown
        - Sección de Advertencias (únicamente si existen)
        - Pie de página discreto
        - Retorna un búfer BytesIO en la posición inicial listo para streaming.
        """
        # Validaciones de negocio
        title = request.title.strip() if request.title else ""
        content = request.content.strip() if request.content else ""

        if not title:
            raise WordExportEmptyContentError("El título no puede estar vacío.")
        if not content:
            raise WordExportEmptyContentError("El contenido a exportar no puede estar vacío.")

        char_len = len(content)
        if char_len > settings.MAX_EXPORT_CHARACTERS:
            raise WordExportSizeExceededError(char_len, settings.MAX_EXPORT_CHARACTERS)

        doc = docx.Document()

        # 1. Configuración de márgenes estándar (1 pulgada / 2.54 cm)
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # 2. Configuración de fuente estándar (Arial)
        normal_style = doc.styles["Normal"]
        normal_style.font.name = self.FONT_NAME
        normal_style.font.size = Pt(11)
        normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

        # 3. Título Principal
        title_p = doc.add_paragraph(style="Title")
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(6)
        title_run = title_p.add_run(title)
        title_run.font.name = self.FONT_NAME
        title_run.font.size = Pt(22)
        title_run.bold = True
        title_run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

        # 4. Metadatos discretos bajo el título (si existen)
        if request.metadata and request.metadata.prompt_name:
            meta_p = doc.add_paragraph()
            meta_p.paragraph_format.space_before = Pt(0)
            meta_p.paragraph_format.space_after = Pt(14)
            meta_run = meta_p.add_run(f"Tipo de análisis: {request.metadata.prompt_name.strip()}")
            meta_run.font.name = self.FONT_NAME
            meta_run.font.size = Pt(10)
            meta_run.italic = True
            meta_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        # 5. Renderizado del contenido principal
        self._render_markdown_content(doc, content)

        # 6. Sección de Advertencias (únicamente si existen elementos)
        if request.warnings and len(request.warnings) > 0:
            warn_heading = doc.add_heading(level=2)
            warn_heading.paragraph_format.space_before = Pt(16)
            warn_heading.paragraph_format.space_after = Pt(6)
            warn_heading_run = warn_heading.add_run("Advertencias")
            warn_heading_run.font.name = self.FONT_NAME
            warn_heading_run.bold = True
            warn_heading_run.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)  # Ámbar oscuro

            for warning in request.warnings:
                if warning and warning.strip():
                    p_warn = doc.add_paragraph(style="List Bullet")
                    self._add_formatted_runs(p_warn, warning.strip())

        # 7. Pie de página nativo en la sección
        section = doc.sections[0]
        footer = section.footer
        p_footer = footer.paragraphs[0]
        p_footer.text = "Generado mediante HITCHINGS"
        p_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if p_footer.runs:
            p_footer.runs[0].font.name = self.FONT_NAME
            p_footer.runs[0].font.size = Pt(9)
            p_footer.runs[0].font.italic = True
            p_footer.runs[0].font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

        # 8. Empaquetar y guardar 100% en memoria
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        logger.info(
            "Documento Word generado exitosamente en memoria (bytes: %d, warnings: %d)",
            buffer.getbuffer().nbytes,
            len(request.warnings),
        )

        return buffer


word_export_service = WordExportService()
